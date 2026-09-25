# -*- coding: utf-8 -*-
"""
把 docs/论文/ 下的五章 Markdown 合并导出为 论文初稿.docx。

设计取舍：
- 用 python-docx 逐块解析 Markdown（标题 / 段落 / 列表 / 引用 / 表格 / 代码块），
  不追求 100% 的 CommonMark 兼容，只覆盖本稿实际用到的语法。
- 中文用「宋体」正文 + 「黑体」标题；显式设置 w:eastAsia，否则 Word 里中文会走默认字体。
- Mermaid 代码块以等宽字体原样保留，并加一行图注说明「源码形式」，
  因为本机没有 mermaid 渲染器（不安装、不猜测替代方案）。
- 表格统一 Table Grid + 表头加粗。
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

BASE = Path(__file__).resolve().parent
CHAPTERS = [
    "01-需求分析.md",
    "02-概要设计.md",
    "03-详细设计.md",
    "04-系统实现.md",
    "05-系统测试.md",
]
OUT = BASE / "论文初稿.docx"

BODY_CJK = "宋体"
HEAD_CJK = "黑体"
MONO = "Consolas"


def set_cjk(run, cjk=BODY_CJK, latin=None):
    run.font.name = latin or cjk
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), cjk)
    rfonts.set(qn("w:ascii"), latin or cjk)
    rfonts.set(qn("w:hAnsi"), latin or cjk)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_inline(par, text, cjk=BODY_CJK, base_bold=False):
    """把 **粗体** 与 `代码` 拆成 run；其余按普通文本。"""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = par.add_run(part[2:-2])
            r.bold = True
            set_cjk(r, cjk)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            r = par.add_run(part[1:-1])
            r.font.name = MONO
            set_cjk(r, MONO, MONO)
            r.font.size = Pt(9.5)
        else:
            r = par.add_run(part)
            r.bold = base_bold
            set_cjk(r, cjk)


def add_code_block(doc, lines, caption=None):
    if caption:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(caption)
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        set_cjk(r)
    for ln in lines:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.left_indent = Cm(0.6)
        pf.line_spacing = 1.0
        r = p.add_run(ln if ln.strip() else " ")
        r.font.name = MONO
        set_cjk(r, MONO, MONO)
        r.font.size = Pt(9)


def split_table_row(line):
    """按未被转义的 | 切分单元格；\\| 视为单元格内的字面竖线。"""
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|") and not s.endswith("\\|"):
        s = s[:-1]
    cells = re.split(r"(?<!\\)\|", s)
    return [c.strip().replace("\\|", "|") for c in cells]


def is_sep_row(line):
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{2,}:?", c or "") for c in cells)


def add_table(doc, rows):
    ncol = max(len(r) for r in rows)
    rows = [r + [""] * (ncol - len(r)) for r in rows]
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = t.cell(i, j)
            c.text = ""
            par = c.paragraphs[0]
            par.paragraph_format.space_before = Pt(1)
            par.paragraph_format.space_after = Pt(1)
            add_inline(par, cell, base_bold=(i == 0))
            for r in par.runs:
                r.font.size = Pt(9)
                if i == 0:
                    r.bold = True
    doc.add_paragraph()


def convert(md_path, doc, first_chapter):
    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        # 代码块
        m = re.match(r"^```(\w*)", stripped)
        if m:
            lang = m.group(1)
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1  # 跳过结束 ```
            if lang == "mermaid":
                add_code_block(doc, buf, caption="［图：Mermaid 源码，渲染图请见 docs/论文/ 下的 Markdown 源文件］")
            else:
                add_code_block(doc, buf)
            continue

        # 表格
        if stripped.startswith("|") and i + 1 < n and is_sep_row(lines[i + 1]):
            rows = [split_table_row(stripped)]
            i += 2
            while i < n and lines[i].strip().startswith("|"):
                rows.append(split_table_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1:
                if not first_chapter:
                    doc.add_page_break()
                p = doc.add_heading(level=0)
                add_inline(p, text, HEAD_CJK)
                for r in p.runs:
                    r.font.size = Pt(22)
            else:
                p = doc.add_heading(level=min(level - 1, 4))
                add_inline(p, text, HEAD_CJK)
            i += 1
            continue

        # 分隔线
        if re.fullmatch(r"-{3,}|\*{3,}", stripped):
            i += 1
            continue

        # 引用块（连续多行合并）
        if stripped.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip()[1:].strip())
                i += 1
            text = " ".join(x for x in buf if x)
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(0.8)
            pf.space_before = Pt(3)
            pf.space_after = Pt(3)
            add_inline(p, text)
            for r in p.runs:
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            continue

        # 无序列表
        m = re.match(r"^(\s*)[-*+]\s+(.*)$", line)
        if m:
            depth = len(m.group(1)) // 2
            p = doc.add_paragraph(style="List Bullet" if depth == 0 else "List Bullet 2")
            add_inline(p, m.group(2))
            i += 1
            continue

        # 有序列表
        m = re.match(r"^(\s*)(\d+)[.)]\s+(.*)$", line)
        if m:
            depth = len(m.group(1)) // 2
            p = doc.add_paragraph(style="List Number" if depth == 0 else "List Number 2")
            add_inline(p, m.group(3))
            i += 1
            continue

        # 空行
        if not stripped:
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(2)
        pf.space_after = Pt(4)
        pf.line_spacing = 1.4
        pf.first_line_indent = Cm(0.74)
        add_inline(p, stripped)
        i += 1


def main():
    doc = Document()

    # 页面与默认样式
    sec = doc.sections[0]
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.8)
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_CJK
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_CJK)

    for idx, name in enumerate(CHAPTERS):
        path = BASE / name
        if not path.exists():
            print(f"[skip] {name} not found", file=sys.stderr)
            continue
        convert(path, doc, first_chapter=(idx == 0))

    doc.save(OUT)
    print(f"saved: {OUT}")


if __name__ == "__main__":
    main()
